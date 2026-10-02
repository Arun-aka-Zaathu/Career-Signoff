# ==================================================
# CareerSignoff Validator Engine
# ==================================================

import os
import openpyxl
import config

# to check if XLS exists
def check_file_exists(file_path):

    if os.path.exists(file_path):
        return True

    return False
def check_file_type(file_path):

    file_extension = os.path.splitext(file_path)[1]

    if file_extension == config.REQUIRED_FILE_TYPE:
        return True

    return False

#to check if the desired xls sheets are present!
def check_required_sheets(file_path):

    workbook = openpyxl.load_workbook(file_path)

    available_sheets = workbook.sheetnames

    missing_sheets = []

    for sheet in config.REQUIRED_SHEETS:

        if sheet not in available_sheets:
            missing_sheets.append(sheet)

    if len(missing_sheets) == 0:
        return True

    return missing_sheets

# column data checking
def check_required_columns(file_path, sheet_name, required_columns):

    workbook = openpyxl.load_workbook(file_path)

    sheet = workbook[sheet_name]

    existing_columns = []

    for cell in sheet[1]:
        existing_columns.append(cell.value)


    missing_columns = []

    for column in required_columns:

        if column not in existing_columns:
            missing_columns.append(column)


    if len(missing_columns) == 0:
        return True

    return missing_columns